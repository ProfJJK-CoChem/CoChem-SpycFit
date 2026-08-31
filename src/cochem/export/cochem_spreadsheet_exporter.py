"""
CoChem Multi-Format Spreadsheet Extraction & Data Interchange Engine (REQ-QC-099).

Provenance:
- Method Matrix [M]: Standardized multi-sheet Excel (.xlsx) workbooks with dynamic formulas
  and RFC 4180 CSV export for computational datasets and machine learning feature extraction.
- OpenXML styling (Dark blue headers #1F497D, white text, alternating row zebra shading #F2F2F2).
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field

from cochem.core.cochem_exceptions import SpreadsheetExportError


class MolecularRecordSpreadsheetRow(BaseModel):
    """Complete molecular calculation descriptor row for spreadsheet export."""
    model_config = ConfigDict(frozen=True)

    job_label: str = Field(description="Unique calculation/species identifier")
    formula: str = Field(description="Stoichiometric formula")
    smiles: Optional[str] = Field(default=None, description="SMILES string representation")
    charge: int = Field(default=0, description="Molecular net charge")
    multiplicity: int = Field(default=1, description="Spin multiplicity (2S+1)")
    e_scf_hartree: float = Field(description="Electronic SCF energy in Hartree")
    zpve_hartree: float = Field(description="Zero-point vibrational energy in Hartree")
    h_298_hartree: float = Field(description="Total enthalpy at 298.15 K in Hartree")
    g_298_hartree: float = Field(description="Total Gibbs free energy at 298.15 K in Hartree")
    dipole_x_debye: float = Field(default=0.0, description="Dipole moment X component in Debye")
    dipole_y_debye: float = Field(default=0.0, description="Dipole moment Y component in Debye")
    dipole_z_debye: float = Field(default=0.0, description="Dipole moment Z component in Debye")
    dipole_total_debye: float = Field(default=0.0, description="Total magnitude of dipole moment in Debye")
    rotational_a_mhz: float = Field(default=0.0, description="Rotational constant A in MHz")
    rotational_b_mhz: float = Field(default=0.0, description="Rotational constant B in MHz")
    rotational_c_mhz: float = Field(default=0.0, description="Rotational constant C in MHz")
    lowest_frequency_cm1: float = Field(default=0.0, description="Lowest harmonic frequency in cm^-1")
    imaginary_frequency_count: int = Field(default=0, description="Number of imaginary frequencies (0 for min, 1 for TS)")
    coordinates_xyz_str: Optional[str] = Field(default=None, description="Cartesian coordinates in XYZ format")


class SpreadsheetExportManifest(BaseModel):
    """Manifest describing exported Excel and CSV spreadsheet files."""
    model_config = ConfigDict(frozen=True)

    total_records: int = Field(description="Number of exported records")
    excel_output_path: str = Field(description="Path to generated .xlsx workbook")
    csv_output_path: str = Field(description="Path to generated .csv data file")
    rows: List[MolecularRecordSpreadsheetRow] = Field(description="Exported record list")


def export_molecular_dataset_to_excel(
    records: List[MolecularRecordSpreadsheetRow],
    output_excel_path: Union[str, Path],
    output_csv_path: Union[str, Path],
) -> SpreadsheetExportManifest:
    """
    Generate professional multi-sheet Excel workbook and RFC 4180 CSV file.

    Args:
        records: List of molecular calculation data rows.
        output_excel_path: Destination path for `.xlsx` workbook.
        output_csv_path: Destination path for `.csv` file.

    Returns:
        SpreadsheetExportManifest: Summary manifest containing paths and record counts.

    Raises:
        SpreadsheetExportError: If workbook building or file writing fails.
    """
    excel_p = Path(output_excel_path)
    csv_p = Path(output_csv_path)

    excel_p.parent.mkdir(parents=True, exist_ok=True)
    csv_p.parent.mkdir(parents=True, exist_ok=True)

    # Styling definitions
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    try:
        wb = openpyxl.Workbook()

        # =========================================================================
        # Sheet 1: Summary & Energetics
        # =========================================================================
        ws1 = wb.active
        ws1.title = "Summary & Energetics"

        headers1 = [
            "Job Label", "Formula", "Charge", "Mult", "E_SCF (Hartree)", "ZPVE (Hartree)",
            "H_298 (Hartree)", "G_298 (Hartree)", "Delta E (kcal/mol)", "Delta G (kcal/mol)", "N_imag"
        ]
        ws1.append(headers1)

        num_rows = len(records)
        last_row = num_rows + 1

        for r_idx, rec in enumerate(records, start=2):
            # Dynamic Excel formulas for relative energies in kcal/mol
            de_formula = f"=($E{r_idx}-MIN($E$2:$E${last_row}))*627.509474"
            dg_formula = f"=($H{r_idx}-MIN($H$2:$H${last_row}))*627.509474"

            row_data = [
                rec.job_label,
                rec.formula,
                rec.charge,
                rec.multiplicity,
                rec.e_scf_hartree,
                rec.zpve_hartree,
                rec.h_298_hartree,
                rec.g_298_hartree,
                de_formula,
                dg_formula,
                rec.imaginary_frequency_count,
            ]
            ws1.append(row_data)

        # Apply formatting to Sheet 1
        for cell in ws1[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx in range(2, last_row + 1):
            is_zebra = (row_idx % 2 == 0)
            for col_idx in range(1, len(headers1) + 1):
                c = ws1.cell(row=row_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                if is_zebra:
                    c.fill = zebra_fill
                # Number formats
                if col_idx in [5, 6, 7, 8]:
                    c.number_format = "0.00000000"
                elif col_idx in [9, 10]:
                    c.number_format = "0.00"
                elif col_idx in [3, 4, 11]:
                    c.number_format = "0"

        # =========================================================================
        # Sheet 2: Dipoles & Spectroscopic
        # =========================================================================
        ws2 = wb.create_sheet(title="Dipoles & Spectroscopic")
        headers2 = [
            "Job Label", "Formula", "mu_x (D)", "mu_y (D)", "mu_z (D)", "mu_total (D)",
            "A (MHz)", "B (MHz)", "C (MHz)", "Ray's Kappa", "Lowest Freq (cm^-1)"
        ]
        ws2.append(headers2)

        for r_idx, rec in enumerate(records, start=2):
            # Ray's asymmetry parameter kappa = (2B - A - C)/(A - C)
            a, b, c_rot = rec.rotational_a_mhz, rec.rotational_b_mhz, rec.rotational_c_mhz
            kappa_formula = f"=(2*$H{r_idx}-$G{r_idx}-$I{r_idx})/($G{r_idx}-$I{r_idx})" if (a - c_rot != 0) else "0.0"

            row_data = [
                rec.job_label,
                rec.formula,
                rec.dipole_x_debye,
                rec.dipole_y_debye,
                rec.dipole_z_debye,
                rec.dipole_total_debye,
                rec.rotational_a_mhz,
                rec.rotational_b_mhz,
                rec.rotational_c_mhz,
                kappa_formula,
                rec.lowest_frequency_cm1,
            ]
            ws2.append(row_data)

        for cell in ws2[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx in range(2, last_row + 1):
            is_zebra = (row_idx % 2 == 0)
            for col_idx in range(1, len(headers2) + 1):
                c = ws2.cell(row=row_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                if is_zebra:
                    c.fill = zebra_fill
                if col_idx in [3, 4, 5, 6]:
                    c.number_format = "0.0000"
                elif col_idx in [7, 8, 9]:
                    c.number_format = "0.00"
                elif col_idx == 10:
                    c.number_format = "0.0000"
                elif col_idx == 11:
                    c.number_format = "0.0"

        # =========================================================================
        # Sheet 3: Cartesian Geometries
        # =========================================================================
        ws3 = wb.create_sheet(title="Cartesian Geometries")
        headers3 = ["Job Label", "Formula", "XYZ Block"]
        ws3.append(headers3)

        for r_idx, rec in enumerate(records, start=2):
            xyz_val = rec.coordinates_xyz_str or f"1\n{rec.job_label}\nX 0.000000 0.000000 0.000000"
            ws3.append([rec.job_label, rec.formula, xyz_val])

        for cell in ws3[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx in range(2, last_row + 1):
            for col_idx in range(1, 4):
                c = ws3.cell(row=row_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border

        # Auto-adjust column widths
        for ws in [ws1, ws2, ws3]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(str(excel_p))

    except Exception as err:
        raise SpreadsheetExportError(
            format_type="Excel OpenXML (.xlsx)",
            reason=f"Failed to build Excel workbook: {str(err)}",
        ) from err

    # =========================================================================
    # RFC 4180 CSV Export
    # =========================================================================
    try:
        with csv_p.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh, delimiter=",", quoting=csv.QUOTE_MINIMAL)
            csv_headers = [
                "job_label", "formula", "smiles", "charge", "multiplicity",
                "e_scf_hartree", "zpve_hartree", "h_298_hartree", "g_298_hartree",
                "dipole_x_debye", "dipole_y_debye", "dipole_z_debye", "dipole_total_debye",
                "rotational_a_mhz", "rotational_b_mhz", "rotational_c_mhz",
                "lowest_frequency_cm1", "imaginary_frequency_count"
            ]
            writer.writerow(csv_headers)
            for rec in records:
                writer.writerow([
                    rec.job_label, rec.formula, rec.smiles or "", rec.charge, rec.multiplicity,
                    f"{rec.e_scf_hartree:.8f}", f"{rec.zpve_hartree:.8f}", f"{rec.h_298_hartree:.8f}", f"{rec.g_298_hartree:.8f}",
                    f"{rec.dipole_x_debye:.4f}", f"{rec.dipole_y_debye:.4f}", f"{rec.dipole_z_debye:.4f}", f"{rec.dipole_total_debye:.4f}",
                    f"{rec.rotational_a_mhz:.2f}", f"{rec.rotational_b_mhz:.2f}", f"{rec.rotational_c_mhz:.2f}",
                    f"{rec.lowest_frequency_cm1:.1f}", rec.imaginary_frequency_count
                ])
    except Exception as err:
        raise SpreadsheetExportError(
            format_type="CSV (RFC 4180)",
            reason=f"Failed to write CSV file: {str(err)}",
        ) from err

    return SpreadsheetExportManifest(
        total_records=len(records),
        excel_output_path=str(excel_p.resolve()),
        csv_output_path=str(csv_p.resolve()),
        rows=records,
    )
